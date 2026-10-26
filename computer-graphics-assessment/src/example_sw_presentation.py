"""예시 작품: 대구스마트고등학교 소프트웨어과 소개 발표 자료 (16:9, 7장).
콘셉트 키워드: 친근함 · 스마트 · 코드
아이콘은 모두 기본 도형·선을 조합해 직접 만든다(노란 원 + 네이비 단색 도형 규칙)."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ASSET = sys.argv[1]
OUT = sys.argv[2]

NAVY = RGBColor(0x0F, 0x2A, 0x4D)     # 교복색 (주조색)
YELLOW = RGBColor(0xFF, 0xC2, 0x0E)   # 부리색 (강조색)
GREEN = RGBColor(0x7A, 0xB8, 0x48)    # 캐릭터 몸 (보조색)
FACE = RGBColor(0xFF, 0xFD, 0xF7)     # 얼굴색 (배경)
PINK = RGBColor(0xFF, 0xC1, 0xC6)     # 볼색
LINE = RGBColor(0x1A, 0x1A, 0x1A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GRAYT = RGBColor(0x5A, 0x64, 0x72)
TITLE_F, BODY_F, CODE_F = 'NanumSquareRound', 'NanumBarunGothic', 'NanumGothicCoding'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def shape(sl, kind, x, y, w, h, fill=None, line=None, lw=0, rot=0, adj=None):
    s = sl.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    if rot:
        s.rotation = rot
    if adj is not None:
        s.adjustments[0] = adj
    return s


def text(sl, x, y, w, h, runs, size=18, font=BODY_F, color=NAVY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """runs: str | list of lines; each line str or list of (text, {opts})"""
    tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = runs if isinstance(runs, list) else [runs]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing
        for seg in (ln if isinstance(ln, list) else [(ln, {})]):
            t, o = seg
            r = p.add_run(); r.text = t
            f = r.font; f.name = o.get('font', font); f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold); f.color.rgb = o.get('color', color)
    return tb


def line(sl, x1, y1, x2, y2, color=NAVY, w=3):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(w)
    return c


def pic(sl, path, x, y, w=None, h=None):
    return sl.shapes.add_picture(f'{ASSET}/{path}', Inches(x), Inches(y),
                                 Inches(w) if w else None, Inches(h) if h else None)


def bg(sl, color):
    f = sl.background.fill; f.solid(); f.fore_color.rgb = color


# ---------- 공통 레이아웃 ----------
def frame(sl, no, title, sub=None):
    bg(sl, FACE)
    # 코드 괄호 모티프 + 제목
    text(sl, 0.6, 0.42, 1.0, 0.7, '</>', size=30, font=CODE_F, color=YELLOW, bold=True)
    text(sl, 1.55, 0.38, 9.5, 0.8, title, size=34, font=TITLE_F, color=NAVY, bold=True)
    if sub:
        text(sl, 1.58, 1.12, 9.5, 0.45, sub, size=17, color=GRAYT)
    # 학교 엠블럼(흰 카드 위)
    shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 12.05, 0.32, 0.8, 0.85, fill=WHITE, line=PINK, lw=1, adj=0.25)
    pic(sl, 'emblem.png', 12.13, 0.38, h=0.72)
    # 하단 띠와 페이지 번호
    shape(sl, MSO_SHAPE.RECTANGLE, 0, 7.1, 13.333, 0.4, fill=NAVY)
    text(sl, 0.6, 7.16, 8, 0.3, [[('대구스마트고등학교  ', {'bold': True, 'color': WHITE}), ('소프트웨어과', {'color': YELLOW, 'bold': True})]], size=13)
    shape(sl, MSO_SHAPE.OVAL, 12.35, 6.88, 0.58, 0.58, fill=YELLOW)
    text(sl, 12.35, 6.88, 0.58, 0.58, str(no), size=18, font=TITLE_F, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def card(sl, x, y, w, h, fill=WHITE, border=PINK):
    return shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=border, lw=1.25, adj=0.12)


# ---------- 아이콘 (노란 원 + 네이비 도형) ----------
def icon_base(sl, cx, cy, r=0.62):
    shape(sl, MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, fill=YELLOW)


def ic_web(sl, cx, cy):
    icon_base(sl, cx, cy)
    shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.38, cy - 0.3, 0.76, 0.6, line=NAVY, lw=3, adj=0.12)
    line(sl, cx - 0.38, cy - 0.12, cx + 0.38, cy - 0.12)
    for i in range(3):
        shape(sl, MSO_SHAPE.OVAL, cx - 0.3 + i * 0.11, cy - 0.24, 0.07, 0.07, fill=NAVY)
    text(sl, cx - 0.3, cy - 0.08, 0.6, 0.35, '</>', size=13, font=CODE_F, bold=True, align=PP_ALIGN.CENTER)


def ic_game(sl, cx, cy):
    icon_base(sl, cx, cy)
    shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.42, cy - 0.22, 0.84, 0.46, fill=NAVY, adj=0.45)
    shape(sl, MSO_SHAPE.RECTANGLE, cx - 0.3, cy - 0.025, 0.22, 0.05, fill=YELLOW)
    shape(sl, MSO_SHAPE.RECTANGLE, cx - 0.215, cy - 0.11, 0.05, 0.22, fill=YELLOW)
    shape(sl, MSO_SHAPE.OVAL, cx + 0.1, cy - 0.1, 0.09, 0.09, fill=YELLOW)
    shape(sl, MSO_SHAPE.OVAL, cx + 0.22, cy + 0.0, 0.09, 0.09, fill=YELLOW)


def ic_ai(sl, cx, cy):
    icon_base(sl, cx, cy)
    for d in (-0.15, 0, 0.15):
        line(sl, cx + d, cy - 0.4, cx + d, cy + 0.4, w=2.5)
        line(sl, cx - 0.4, cy + d, cx + 0.4, cy + d, w=2.5)
    shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.28, cy - 0.28, 0.56, 0.56, fill=NAVY, adj=0.15)
    text(sl, cx - 0.28, cy - 0.28, 0.56, 0.56, 'AI', size=15, font=TITLE_F, bold=True, color=YELLOW,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def ic_data(sl, cx, cy):
    icon_base(sl, cx, cy)
    for i, hh in enumerate((0.25, 0.45, 0.62)):
        shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.3 + i * 0.22, cy + 0.3 - hh, 0.15, hh, fill=NAVY, adj=0.2)
    line(sl, cx - 0.38, cy + 0.33, cx + 0.38, cy + 0.33)


def ic_badge(sl, cx, cy):
    icon_base(sl, cx, cy)
    shape(sl, MSO_SHAPE.RECTANGLE, cx - 0.2, cy + 0.02, 0.12, 0.38, fill=NAVY, rot=15)
    shape(sl, MSO_SHAPE.RECTANGLE, cx + 0.08, cy + 0.02, 0.12, 0.38, fill=NAVY, rot=-15)
    shape(sl, MSO_SHAPE.OVAL, cx - 0.27, cy - 0.38, 0.54, 0.54, fill=NAVY)
    shape(sl, MSO_SHAPE.OVAL, cx - 0.16, cy - 0.27, 0.32, 0.32, line=YELLOW, lw=2.5)


def ic_monitor(sl, cx, cy):
    icon_base(sl, cx, cy)
    shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.4, cy - 0.32, 0.8, 0.5, fill=NAVY, adj=0.1)
    text(sl, cx - 0.4, cy - 0.32, 0.8, 0.5, '</>', size=14, font=CODE_F, bold=True, color=YELLOW,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    shape(sl, MSO_SHAPE.RECTANGLE, cx - 0.05, cy + 0.18, 0.1, 0.12, fill=NAVY)
    shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.22, cy + 0.29, 0.44, 0.07, fill=NAVY, adj=0.5)


def ic_network(sl, cx, cy):
    icon_base(sl, cx, cy)
    pts = [(cx, cy - 0.28), (cx - 0.3, cy + 0.22), (cx + 0.3, cy + 0.22)]
    for (a, b) in ((0, 1), (1, 2), (0, 2)):
        line(sl, *pts[a], *pts[b])
    for (x, y) in pts:
        shape(sl, MSO_SHAPE.OVAL, x - 0.13, y - 0.13, 0.26, 0.26, fill=NAVY)


def ic_building(sl, cx, cy):
    icon_base(sl, cx, cy)
    shape(sl, MSO_SHAPE.ISOSCELES_TRIANGLE, cx - 0.42, cy - 0.4, 0.84, 0.25, fill=NAVY)
    shape(sl, MSO_SHAPE.RECTANGLE, cx - 0.36, cy + 0.26, 0.72, 0.08, fill=NAVY)
    for i in range(4):
        shape(sl, MSO_SHAPE.RECTANGLE, cx - 0.3 + i * 0.18, cy - 0.12, 0.08, 0.36, fill=NAVY)


def ic_cap(sl, cx, cy):
    icon_base(sl, cx, cy)
    shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx - 0.22, cy - 0.02, 0.44, 0.26, fill=NAVY, adj=0.3)
    shape(sl, MSO_SHAPE.DIAMOND, cx - 0.46, cy - 0.32, 0.92, 0.42, fill=NAVY)
    line(sl, cx + 0.3, cy - 0.11, cx + 0.3, cy + 0.2, color=WHITE, w=2)
    shape(sl, MSO_SHAPE.OVAL, cx + 0.255, cy + 0.17, 0.09, 0.09, fill=WHITE)


# ===== 1. 표지 =====
s = prs.slides.add_slide(BLANK); bg(s, NAVY)
text(s, 7.2, -0.4, 6, 4, '</>', size=220, font=CODE_F, color=RGBColor(0x1C, 0x3D, 0x66), bold=True)
shape(s, MSO_SHAPE.OVAL, 8.45, 1.15, 4.6, 4.6, fill=YELLOW)
pic(s, 'mascot_front.png', 9.25, 1.0, h=5.9)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, 0.75, 0.95, 1.0, fill=WHITE, adj=0.2)
pic(s, 'emblem.png', 0.88, 0.82, h=0.86)
text(s, 1.95, 0.88, 6, 0.4, '대구스마트고등학교', size=22, font=TITLE_F, color=WHITE, bold=True)
text(s, 1.95, 1.3, 6, 0.4, '2027학년도 신입생 학과 소개', size=16, color=PINK)
text(s, 0.8, 2.45, 8, 0.6, '아이디어를 현실로!', size=34, font=TITLE_F, color=YELLOW, bold=True)
text(s, 0.75, 2.95, 8, 1.6, '소프트웨어과', size=88, font=TITLE_F, color=WHITE, bold=True)
shape(s, MSO_SHAPE.RECTANGLE, 0.85, 4.85, 1.2, 0.09, fill=YELLOW)
text(s, 0.8, 5.15, 7.5, 1.0, ['세상을 바꾸는 코드,', '인공지능과 함께 시작하세요!'], size=22, color=WHITE)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, 6.45, 3.9, 0.55, fill=GREEN, adj=0.5)
text(s, 0.8, 6.45, 3.9, 0.55, '안내 캐릭터 · 스마영', size=17, font=TITLE_F, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ===== 2. 우리 학교 =====
s = prs.slides.add_slide(BLANK)
frame(s, 2, '대구스마트고등학교를 소개해요', '기술로 꿈을 키우는 남녀 공학 특성화고등학교')
p = pic(s, 'building_wide.jpg', 0.6, 1.75, w=9.0)
p.line.color.rgb = NAVY; p.line.width = Pt(2)
text(s, 0.6, 5.75, 3, 0.4, '모집 학과', size=18, font=TITLE_F, bold=True)
for i, (name, on) in enumerate((('정밀기계과', 0), ('자동화기계과', 0), ('소프트웨어과', 1))):
    x = 0.6 + i * 3.05
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 6.2, 2.85, 0.65, fill=YELLOW if on else WHITE, line=NAVY, lw=1.5, adj=0.5)
    text(s, x, 6.2, 2.85, 0.65, name, size=20, font=TITLE_F, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
shape(s, MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, 9.95, 1.85, 2.95, 1.75, fill=WHITE, line=NAVY, lw=1.5)
text(s, 10.15, 2.0, 2.6, 1.4, ['안녕! 나는 스마영이야.', [('소프트웨어과', {'bold': True, 'color': GREEN}), ('를', {})], '소개할게!'],
     size=18, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
s.shapes[-2].adjustments[0], s.shapes[-2].adjustments[1] = -0.1, 0.75
pic(s, 'mascot_thumb.png', 10.35, 4.05, h=2.85)

# ===== 3. 교육목표 =====
s = prs.slides.add_slide(BLANK)
frame(s, 3, '소프트웨어과는 이런 곳이에요', '학과 교육목표')
p = pic(s, 'sw_lab.jpg', 0.6, 1.85, w=6.2)
p.line.color.rgb = NAVY; p.line.width = Pt(2)
text(s, 0.6, 4.85, 6.2, 0.4, '▲ 소프트웨어과 실습 모습', size=15, color=GRAYT)
card(s, 7.2, 1.85, 5.55, 4.6, fill=NAVY, border=NAVY)
text(s, 7.6, 2.05, 2, 0.9, '“', size=72, font=TITLE_F, color=YELLOW, bold=True)
text(s, 7.6, 2.95, 4.8, 2.6, [
    [('앱·웹·게임 개발', {'color': YELLOW, 'bold': True}), (',', {})],
    [('인공지능', {'color': YELLOW, 'bold': True}), (', ', {}), ('빅데이터', {'color': YELLOW, 'bold': True}), (' 분야의', {})],
    '실무 역량을 갖춘',
    '소프트웨어 인재 양성'], size=26, font=TITLE_F, color=WHITE, bold=True, spacing=1.3)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 5.5, 6.2, 0.95, fill=WHITE, line=PINK, lw=1.25, adj=0.2)
text(s, 0.85, 5.5, 5.8, 0.95, [[('아이디어를 현실로!', {'bold': True, 'color': GREEN}), (' 생각한 것을 코드로 직접 만들어 봐요.', {})]],
     size=18, anchor=MSO_ANCHOR.MIDDLE)

# ===== 4. 배우는 분야 =====
s = prs.slides.add_slide(BLANK)
frame(s, 4, '이런 분야를 배워요', '교육목표의 네 가지 핵심 분야')
items = [(ic_web, '앱·웹 개발', ['웹사이트와 모바일 앱', '서비스를 만드는 기술']),
         (ic_game, '게임 개발', ['게임을 기획하고', '프로그래밍하는 기술']),
         (ic_ai, '인공지능', ['AI의 원리를 이해하고', '활용하는 기술']),
         (ic_data, '빅데이터', ['데이터를 모으고', '분석하는 기술'])]
for i, (fn, name, desc) in enumerate(items):
    x = 0.6 + i * 3.08
    card(s, x, 1.9, 2.85, 4.75)
    fn(s, x + 1.425, 2.95)
    text(s, x, 3.85, 2.85, 0.6, name, size=26, font=TITLE_F, bold=True, align=PP_ALIGN.CENTER)
    shape(s, MSO_SHAPE.RECTANGLE, x + 1.175, 4.55, 0.5, 0.06, fill=YELLOW)
    text(s, x + 0.1, 4.8, 2.65, 1.2, desc, size=16, align=PP_ALIGN.CENTER, color=GRAYT)

# ===== 5. 자격증 =====
s = prs.slides.add_slide(BLANK)
frame(s, 5, '취득 가능 자격증', '배운 실력을 자격증으로 증명해요')
certs = [('정보처리', ['정보처리산업기사', '정보처리기능사']), ('정보기기', ['정보기기운용기능사']), ('그래픽·사무·네트워크', ['GTQ', 'ITQ', '네트워크관리사'])]
for i, (head, names) in enumerate(certs):
    x = 0.6 + i * 4.1
    card(s, x, 1.9, 3.85, 4.5)
    ic_badge(s, x + 1.925, 2.95)
    text(s, x, 3.85, 3.85, 0.45, head, size=17, color=GRAYT, align=PP_ALIGN.CENTER)
    text(s, x + 0.2, 4.35, 3.45, 1.9, names, size=24, font=TITLE_F, bold=True, align=PP_ALIGN.CENTER, spacing=1.35)
text(s, 0.6, 6.5, 11, 0.4, '※ 이 밖에도 다양한 자격증에 도전할 수 있어요.', size=15, color=GRAYT)

# ===== 6. 진로 =====
s = prs.slides.add_slide(BLANK)
frame(s, 6, '졸업 후에는 이렇게 나아가요', '취업과 진학, 두 갈래 길')
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 1.8, 1.6, 0.55, fill=NAVY, adj=0.5)
text(s, 0.6, 1.8, 1.6, 0.55, '취업', size=20, font=TITLE_F, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
jobs = [(ic_monitor, ['응용프로그램', '웹·앱·게임 개발자']), (ic_network, ['AI·빅데이터', '네트워크 관련 업체']), (ic_building, ['소프트웨어 전문기업', '공기업 전산직'])]
for i, (fn, desc) in enumerate(jobs):
    x = 0.6 + i * 2.95
    card(s, x, 2.55, 2.75, 3.6)
    fn(s, x + 1.375, 3.55)
    text(s, x + 0.1, 4.5, 2.55, 1.3, desc, size=19, font=TITLE_F, bold=True, align=PP_ALIGN.CENTER, spacing=1.3)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 9.65, 1.8, 1.6, 0.55, fill=GREEN, adj=0.5)
text(s, 9.65, 1.8, 1.6, 0.55, '진학', size=20, font=TITLE_F, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
card(s, 9.65, 2.55, 3.1, 3.6, fill=WHITE, border=GREEN)
ic_cap(s, 11.2, 3.55)
text(s, 9.75, 4.5, 2.9, 1.5, [[('대학 진학', {'bold': True, 'font': TITLE_F, 'size': 21})], '졸업생 진학 사례', [('서강대학교', {'bold': True, 'color': GREEN})]],
     size=17, align=PP_ALIGN.CENTER, spacing=1.3)
text(s, 0.6, 6.35, 12, 0.5, [[('취업', {'bold': True}), ('도 ', {}), ('대학', {'bold': True}), ('도, 소프트웨어과에서 시작해요!', {})]], size=20, font=TITLE_F, color=NAVY)

# ===== 7. 모집·마무리 =====
s = prs.slides.add_slide(BLANK); bg(s, NAVY)
text(s, 0.6, 0.42, 1.0, 0.7, '</>', size=30, font=CODE_F, color=YELLOW, bold=True)
text(s, 1.55, 0.38, 10, 0.8, '함께 코드를 만들어 갈 친구를 기다려요!', size=34, font=TITLE_F, color=WHITE, bold=True)
card(s, 0.6, 1.55, 5.4, 4.3, fill=WHITE, border=WHITE)
text(s, 0.9, 1.8, 4.8, 0.5, '2027학년도 모집 인원 (남녀 공학)', size=18, font=TITLE_F, bold=True, color=GRAYT)
text(s, 0.9, 2.3, 4.8, 1.2, [[('1', {'size': 60}), ('학급 ', {'size': 28}), ('24', {'size': 60, 'color': GREEN}), ('명', {'size': 28})]],
     font=TITLE_F, bold=True)
line(s, 0.9, 3.7, 5.7, 3.7, color=PINK, w=1.5)
text(s, 0.9, 3.85, 4.8, 1.2, ['취업희망자전형  16명', '일반전형 및 기타 특별전형  8명'], size=19, spacing=1.35)
text(s, 0.9, 5.05, 4.9, 0.7, '※ 모집 정원은 「2027학년도 대구광역시 고등학교 학생 배치계획」에 따라 변경될 수 있음', size=12, color=GRAYT)
card(s, 6.3, 1.55, 6.45, 4.3, fill=RGBColor(0x1C, 0x3D, 0x66), border=YELLOW)
text(s, 6.65, 1.8, 6, 0.6, [[('OPEN DAY ', {'color': YELLOW}), ('학교설명회', {'color': WHITE})]], size=28, font=TITLE_F, bold=True)
for i, (n, dt, tm) in enumerate((('1차', '9월 19일(토)', '오후 2시 30분'), ('2차', '11월 20일(금)', '오후 1시 30분'), ('3차', '11월 21일(토)', '오후 2시'))):
    y = 2.65 + i * 0.95
    shape(s, MSO_SHAPE.OVAL, 6.65, y, 0.75, 0.75, fill=YELLOW)
    text(s, 6.65, y, 0.75, 0.75, n, size=17, font=TITLE_F, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 7.55, y, 5.15, 0.75, [[(dt + '  ', {'color': WHITE}), (tm, {'color': YELLOW})]], size=20, font=TITLE_F, bold=True, anchor=MSO_ANCHOR.MIDDLE)
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.6, 6.2, 9.6, 0.85, fill=YELLOW, adj=0.5)
text(s, 0.6, 6.2, 9.6, 0.85, [[('입학 문의  ', {}), ('대구스마트고등학교 ☎ 053-○○○-○○○○', {'bold': True})]], size=24, font=TITLE_F,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
pic(s, 'mascot_thumb.png', 11.2, 5.95, h=1.5)

prs.save(OUT)
print('saved', OUT)
